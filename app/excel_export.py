"""Excel export for analysis results (F5): one sheet per analysis plus a
Summary sheet, via openpyxl per CLAUDE.md (no other spreadsheet lib).
"""
from __future__ import annotations

from dataclasses import dataclass
from typing import Dict, Optional

import numpy as np
from openpyxl import Workbook

from app.project import Project


@dataclass
class AnalysisSeries:
    label: str
    x_label: str
    y_label: str
    x: np.ndarray
    y: np.ndarray


@dataclass
class AnalysisMatrix:
    label: str
    matrix: np.ndarray


def export_to_excel(
    project: Project,
    series: Dict[str, AnalysisSeries],
    matrices: Optional[Dict[str, AnalysisMatrix]] = None,
    output_path: str = "analysis.xlsx",
) -> str:
    matrices = matrices or {}
    workbook = Workbook()

    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Project", project.name])
    summary.append(["Model", project.model_type.value])
    summary.append(["Temperature (K)", project.parameters.temperature_k])
    summary.append(["Steps", project.parameters.n_steps])
    summary.append(["Timestep (fs)", project.parameters.timestep_fs])
    summary.append([])
    summary.append(["Analysis", "Rows"])
    for key, s in series.items():
        summary.append([s.label, len(s.x)])
    for key, m in matrices.items():
        summary.append([m.label, f"{m.matrix.shape[0]}x{m.matrix.shape[1]}"])

    for key, s in series.items():
        sheet = workbook.create_sheet(title=_safe_sheet_title(s.label))
        sheet.append([s.x_label, s.y_label])
        for x_val, y_val in zip(s.x.tolist(), s.y.tolist()):
            sheet.append([x_val, y_val])

    for key, m in matrices.items():
        sheet = workbook.create_sheet(title=_safe_sheet_title(m.label))
        for row in m.matrix.tolist():
            sheet.append(row)

    workbook.save(output_path)
    return output_path


def _safe_sheet_title(label: str) -> str:
    # Excel sheet names: max 31 chars, no []:*?/\\
    cleaned = "".join(c for c in label if c not in '[]:*?/\\')
    return cleaned[:31] or "Sheet"
