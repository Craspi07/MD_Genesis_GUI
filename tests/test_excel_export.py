from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from app.excel_export import AnalysisSeries, AnalysisMatrix, export_to_excel
from app.project import Project


def test_export_to_excel_creates_summary_and_series_sheets(tmp_path: Path):
    project = Project(name="myproj", directory=str(tmp_path))
    series = {
        "rmsd": AnalysisSeries(
            label="RMSD (vs. initial frame)",
            x_label="Frame",
            y_label="RMSD (nm)",
            x=np.array([0, 1, 2]),
            y=np.array([0.0, 0.1, 0.15]),
        )
    }
    matrices = {
        "contact_map": AnalysisMatrix(label="Contact map", matrix=np.array([[1, 0], [0, 1]])),
    }

    output_path = str(tmp_path / "analysis.xlsx")
    export_to_excel(project, series, matrices, output_path)

    workbook = load_workbook(output_path)
    assert "Summary" in workbook.sheetnames
    assert "RMSD (vs. initial frame)" in workbook.sheetnames
    assert "Contact map" in workbook.sheetnames

    rmsd_sheet = workbook["RMSD (vs. initial frame)"]
    assert rmsd_sheet["A1"].value == "Frame"
    assert rmsd_sheet["B1"].value == "RMSD (nm)"
    assert rmsd_sheet["A2"].value == 0
    assert rmsd_sheet["B3"].value == 0.1

    summary_sheet = workbook["Summary"]
    assert summary_sheet["A1"].value == "Project"
    assert summary_sheet["B1"].value == "myproj"


def test_sheet_title_sanitized_and_truncated(tmp_path: Path):
    project = Project(name="p", directory=str(tmp_path))
    long_label = "A" * 40 + "[bad]chars*here"
    series = {
        "x": AnalysisSeries(label=long_label, x_label="x", y_label="y", x=np.array([0]), y=np.array([1.0])),
    }
    output_path = str(tmp_path / "out.xlsx")
    export_to_excel(project, series, {}, output_path)
    workbook = load_workbook(output_path)
    matching = [name for name in workbook.sheetnames if name != "Summary"]
    assert len(matching) == 1
    assert len(matching[0]) <= 31
    assert "[" not in matching[0] and "*" not in matching[0]
